"""
Example: Generate 2Penguins client proposal from template.
"""

from docx import Document
from datetime import date

def generate_2penguins_proposal(client_name: str, project_name: str, budget: float):
    """Generate branded proposal for 2Penguins."""

    # Load template
    template_path = '../ln-800-document-automation/templates/docx/2penguins_proposal.docx'
    doc = Document(template_path)

    # Variable substitution
    variables = {
        'client_name': client_name,
        'project_name': project_name,
        'project_budget': f"€{budget:,.2f}",
        'deadline': '2026-06-30',
        'date': date.today().isoformat()
    }

    # Replace placeholders
    for paragraph in doc.paragraphs:
        for key, value in variables.items():
            placeholder = f"{{{{{key}}}}}"
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, str(value))

    # Save output
    output_path = f'proposals/{client_name.lower().replace(" ", "_")}_proposal.docx'
    doc.save(output_path)

    print(f"✅ Proposal created: {output_path}")
    return output_path


if __name__ == '__main__':
    generate_2penguins_proposal('ACME Corp', 'Digital Signage System', 45000)
